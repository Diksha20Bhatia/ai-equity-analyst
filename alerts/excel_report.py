"""
excel_report.py  —  "The Ledger"
====================================
Writes each swing run's results into a persistent Excel workbook
(output/portfolio_tracker.xlsx) — created on the first run, appended to
(as a new sheet) on every run after that. Re-running on the SAME day
replaces that day's sheet rather than duplicating it, matching how
output/analysis_DATE.json already behaves.

Each date's sheet holds two tables:
  1. Selected Stocks — the deduped final shortlist across every risk tier
     you ran today, with a Tier(s) column showing which tier(s) picked it.
  2. Top N Watchlist — every stock the Scout flagged as unusual today
     (real anomaly score), whether or not it made the final cut.

Plain code, zero AI — every number here was already computed earlier in
the run. The longitudinal price-tracking (how these picks have moved
since they were listed) is a separate, deliberately independent tool —
see tracker/price_tracker.py.
"""

import datetime as dt
from pathlib import Path

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "output" / "portfolio_tracker.xlsx"


def _sheet_name(date: dt.date) -> str:
    return date.isoformat()  # e.g. "2026-08-26" — well under Excel's 31-char sheet-name limit


def write_run(results_by_tier: dict, analyses: list, run_date: dt.date = None) -> str:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    run_date = run_date or dt.date.today()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    if OUTPUT_PATH.exists():
        wb = load_workbook(OUTPUT_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)  # drop the default blank sheet

    sheet_name = _sheet_name(run_date)
    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])  # same-day re-run: replace, don't duplicate
    ws = wb.create_sheet(title=sheet_name)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2A78D6")
    title_font = Font(bold=True, size=12)

    # --- Table 1: selected stocks, deduped across every risk tier ------
    picks_by_symbol = {}
    for tier, result in results_by_tier.items():
        for p in result.get("shortlist", []):
            entry = picks_by_symbol.setdefault(p["symbol"], {
                "current_price": p.get("current_price"),
                "tiers": [],
            })
            entry["tiers"].append(tier)

    sector_by_symbol = {a["symbol"]: a.get("sector") for a in analyses}

    row = 1
    ws.cell(row=row, column=1, value=f"Selected Stocks — {sheet_name}").font = title_font
    row += 1
    for col, h in enumerate(["Symbol", "Description (Sector)", "Current Price (₹)", "Tier(s)"], start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font, c.fill = header_font, header_fill
    row += 1
    for sym, entry in picks_by_symbol.items():
        ws.cell(row=row, column=1, value=sym)
        ws.cell(row=row, column=2, value=sector_by_symbol.get(sym) or "n/a")
        ws.cell(row=row, column=3, value=entry["current_price"])
        ws.cell(row=row, column=4, value=", ".join(entry["tiers"]))
        row += 1

    row += 3  # a few blank spacer rows before the second table

    # --- Table 2: the full watchlist the Scout flagged today -----------
    ws.cell(row=row, column=1, value=f"Top {len(analyses)} Watchlist — {sheet_name}").font = title_font
    row += 1
    for col, h in enumerate(["Symbol", "Current Price (₹)"], start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font, c.fill = header_font, header_fill
    row += 1
    for a in analyses:
        ws.cell(row=row, column=1, value=a["symbol"])
        ws.cell(row=row, column=2, value=a.get("market_stats", {}).get("current_price"))
        row += 1

    for col, width in zip("ABCD", (14, 26, 18, 24)):
        ws.column_dimensions[col].width = width

    wb.save(OUTPUT_PATH)
    print(f"[excel] {sheet_name}: {len(picks_by_symbol)} selected, {len(analyses)} watchlist "
          f"-> {OUTPUT_PATH}")
    return str(OUTPUT_PATH)
