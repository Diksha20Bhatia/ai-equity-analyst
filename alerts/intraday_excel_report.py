"""
intraday_excel_report.py  —  "The Ledger", intraday edition
====================================
Writes each intraday scan's picks into its own persistent workbook
(output/intraday_tracker.xlsx) — a deliberately separate file from
output/portfolio_tracker.xlsx, since intraday picks (entry/stop/target,
breakout direction, pattern tag) don't share a shape with swing picks.
Created on the first run, appended to (as a new sheet) on every run
after that. Re-running on the SAME day replaces that day's sheet rather
than duplicating it — the exact same rule alerts/excel_report.py uses,
so an intraday scan re-run mid-morning updates today's record instead of
piling up one sheet per scan.

The sheet's price column is deliberately named "Current Price (₹)",
matching alerts/excel_report.py's own column exactly — that's what lets
tracker/price_tracker.py track this workbook with zero extra code: it
finds any table whose header row starts with "Symbol" and has a
"Current Price"-labelled column, regardless of which workbook it's in.

Plain code, zero AI — every number here was already computed earlier in
the run.
"""

import datetime as dt
from pathlib import Path

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "output" / "intraday_tracker.xlsx"


def _sheet_name(date: dt.date) -> str:
    return date.isoformat()


def write_run(picks: list, run_date: dt.date = None) -> str:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    run_date = run_date or dt.date.today()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    if OUTPUT_PATH.exists():
        wb = load_workbook(OUTPUT_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)  # drop the default blank sheet

    sheet_name = _sheet_name(run_date)
    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])  # same-day re-run: replace, don't duplicate
    ws = wb.create_sheet(title=sheet_name)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2A78D6")
    title_font = Font(bold=True, size=12)

    row = 1
    ws.cell(row=row, column=1, value=f"Intraday Picks — {sheet_name}").font = title_font
    row += 1
    headers = ["Symbol", "Direction", "Conviction", "Current Price (₹)", "Stop Loss (₹)", "Target (₹)", "Pattern"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font, c.fill = header_font, header_fill
    row += 1
    for p in picks:
        ws.cell(row=row, column=1, value=p["symbol"])
        ws.cell(row=row, column=2, value=p.get("direction"))
        ws.cell(row=row, column=3, value=p.get("conviction"))
        ws.cell(row=row, column=4, value=p.get("entry"))
        ws.cell(row=row, column=5, value=p.get("stop_loss"))
        ws.cell(row=row, column=6, value=p.get("target"))
        ws.cell(row=row, column=7, value=(p.get("pattern_tag") or "").replace("_", " "))
        row += 1

    for col, width in zip("ABCDEFG", (14, 12, 16, 18, 16, 16, 20)):
        ws.column_dimensions[col].width = width

    wb.save(OUTPUT_PATH)
    print(f"[excel] {sheet_name}: {len(picks)} intraday pick(s) -> {OUTPUT_PATH}")
    return str(OUTPUT_PATH)
