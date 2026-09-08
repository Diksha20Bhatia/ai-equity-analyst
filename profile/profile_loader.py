"""
profile_loader.py  —  "Your preferences, in any format"
=======================================================
Reads an investor profile from .txt / .md / .pdf / .html / .docx / .xlsx,
then uses Bodhi's LLM Gateway to normalise the messy text into ONE clean
structured object the Decision Agent can act on. If that call fails, it
does a light keyword-based extraction instead.

The raw text is the one real untrusted-input boundary in this app (see
safety/text_safety.py) — scanned for prompt-injection patterns and
PII-redacted before it ever reaches a prompt.

A .json profile is treated as ALREADY structured (e.g. built from the
dashboard's dropdown form) and loaded directly — no AI call needed,
which is also a free win for the app's token budget.

Returned shape:
  {
    "risk_appetite": "...",
    "capital_available": "...",
    "capital_numeric": 500000 or null,      # INR, only if a real number was stated
    "max_position_pct": 15 or null,         # only if a real max-position rule was stated
    "investment_horizon": "...",
    "sector_preferences": [...],
    "sector_exclusions": [...],
    "constraints": "..."
  }
"""

import os
import json
from config import call_llm
from safety.text_safety import sanitize_untrusted_text


# ---------------------------------------------------------------------
# Step 1: read raw text out of whatever file format was given.
# ---------------------------------------------------------------------
def _read_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            print("[profile] pypdf not installed; cannot read PDF.")
            return ""

    if ext in (".html", ".htm"):
        try:
            from bs4 import BeautifulSoup
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return BeautifulSoup(f.read(), "html.parser").get_text(" ")
        except ImportError:
            print("[profile] beautifulsoup4 not installed; cannot read HTML.")
            return ""

    if ext == ".docx":
        try:
            import docx
            d = docx.Document(path)
            return "\n".join(p.text for p in d.paragraphs)
        except ImportError:
            print("[profile] python-docx not installed; cannot read DOCX.")
            return ""

    if ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except ImportError:
            print("[profile] openpyxl not installed; cannot read XLSX.")
            return ""

    print(f"[profile] unsupported format: {ext}")
    return ""


# ---------------------------------------------------------------------
# Step 2: turn raw text into a structured profile object.
# ---------------------------------------------------------------------
def _structure_with_llm(raw: str) -> dict:
    if not raw.strip():
        return {}
    try:
        prompt = (
            "Read this investor profile and extract the fields below. "
            "If a field is missing, use a sensible default or empty list. "
            "capital_numeric and max_position_pct must be null unless the text states "
            "an actual number — never estimate or guess them.\n\n"
            f"PROFILE TEXT:\n{raw}\n\n"
            'Reply as JSON only: {"risk_appetite": "conservative/moderate/aggressive", '
            '"capital_available": "string", "capital_numeric": <INR number or null>, '
            '"max_position_pct": <number 0-100 or null>, '
            '"investment_horizon": "short/medium/long", '
            '"sector_preferences": [..], "sector_exclusions": [..], "constraints": "string"}'
        )
        text, _usage = call_llm(prompt=prompt, temperature=0.1, json_mode=True)
        text = text.strip().strip("`")
        if "{" in text:
            text = text[text.find("{"):]
        return json.loads(text)
    except Exception as e:  # noqa: BLE001
        print(f"[profile] LLM structuring failed, using keyword fallback. ({e})")
        return {}


def _structure_with_keywords(raw: str) -> dict:
    low = raw.lower()
    risk = "moderate"
    if "aggressive" in low:
        risk = "aggressive"
    elif "conservative" in low or "low risk" in low:
        risk = "conservative"
    horizon = "long" if "long" in low else "medium"
    exclusions = []
    for sector in ("tobacco", "alcohol", "penny", "f&o", "derivatives", "smallcap"):
        if sector in low:
            exclusions.append(sector)
    return {
        "risk_appetite": risk,
        "capital_available": "",
        "capital_numeric": None,   # keyword scan can't reliably extract a real number
        "max_position_pct": None,
        "investment_horizon": horizon,
        "sector_preferences": [],
        "sector_exclusions": exclusions,
        "constraints": "",
    }


# ---------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------
def load_profile(path: str) -> dict:
    if not path or not os.path.exists(path):
        print("[profile] no profile file found — running without personalisation.")
        return {}

    if path.lower().endswith(".json"):
        with open(path, "r", encoding="utf-8") as f:
            structured = json.load(f)
        risk_label = "/".join(structured.get("risk_appetites") or [structured.get("risk_appetite", "?")])
        print(f"[profile] loaded pre-structured profile ({risk_label} risk) "
              "— no AI call needed.")
        return structured

    raw = _read_text(path)
    if not raw.strip():
        return {}
    # This is the one real untrusted-input boundary in the whole app — an
    # uploaded document, arbitrary text from outside the system — so it's
    # the one place a guardrail belongs: logs (doesn't block) anything
    # that looks like a prompt-injection attempt, and redacts obvious PII
    # before it ever reaches a prompt. See safety/text_safety.py.
    raw = sanitize_untrusted_text(raw, source=f"investor_profile:{os.path.basename(path)}")
    structured = _structure_with_llm(raw)
    if not structured:
        structured = _structure_with_keywords(raw)
    print(f"[profile] loaded profile ({structured.get('risk_appetite','?')} risk).")
    return structured
