"""
quality_report.py  —  does the shortlist actually get better over time?
=====================================================================
    uv run python tracker/quality_report.py

Reads the real, already-collected tracking history in
output/portfolio_tracker.xlsx (written by alerts/excel_report.py,
updated by tracker/price_tracker.py) and computes real, honest
statistics: of every stock ever selected into a "Selected Stocks" table,
how many are up vs down since the day they were picked, by how much on
average, and broken down by risk tier — a real evidence trail for
whether this app's picks are actually any good over time, not just a
demo shown once.

Deliberately looks only at "Selected Stocks" tables (identified by their
Tier(s) column) — not the "Top N Watchlist" table (flagged but never
acted on) and not the intraday workbook (same-day calls aren't
meaningfully evaluated on a "since listing" basis the way swing picks
are).

Zero AI, zero new network calls — every number here is a subtraction
over real prices tracker/price_tracker.py already fetched.
"""

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tracker.price_tracker import (  # noqa: E402
    WORKBOOK_PATHS, _find_tables, _reference_price_col, _parse_price,
)

OUT_PATH = ROOT / "output" / "quality_report.json"


def _original_price_col(ws, table: dict):
    """The FIRST real price checkpoint in this table — the original
    listing price, not the most recent tracked column — for a true
    since-day-one comparison."""
    header_row = table["header_row"]
    for col in range(1, table["last_col"] + 1):
        header = str(ws.cell(row=header_row, column=col).value or "")
        if "Current Price" in header:
            return col
    return None


def build_report() -> dict:
    from openpyxl import load_workbook

    picks = []

    for path in WORKBOOK_PATHS:
        if not path.exists():
            continue
        wb = load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for table in _find_tables(ws):
                header_row = table["header_row"]
                headers = [
                    str(ws.cell(row=header_row, column=c).value or "")
                    for c in range(1, table["last_col"] + 1)
                ]
                if "Tier(s)" not in headers:
                    continue  # only "Selected Stocks" tables have this column
                tier_col = headers.index("Tier(s)") + 1
                orig_col = _original_price_col(ws, table)
                latest_col = _reference_price_col(ws, table)
                if not orig_col or not latest_col or latest_col == orig_col:
                    continue  # no tracking history yet for this sheet

                for row in range(table["data_start"], table["data_end"] + 1):
                    symbol = ws.cell(row=row, column=1).value
                    if not symbol:
                        continue
                    orig_price = _parse_price(ws.cell(row=row, column=orig_col).value)
                    latest_price = _parse_price(ws.cell(row=row, column=latest_col).value)
                    tiers = ws.cell(row=row, column=tier_col).value or ""
                    if not orig_price or not latest_price:
                        continue
                    pct = (latest_price - orig_price) / orig_price * 100
                    picks.append({
                        "workbook": path.name, "sheet": sheet_name, "symbol": symbol,
                        "tiers": tiers, "listed_price": orig_price, "latest_price": latest_price,
                        "change_pct": round(pct, 2),
                    })

    if not picks:
        return {
            "total_tracked_picks": 0,
            "summary": "No tracked history yet — run tracker/price_tracker.py at "
                       "least once after a swing run, so there's a second price "
                       "checkpoint to compare against the listing price.",
        }

    up = [p for p in picks if p["change_pct"] > 0]
    down = [p for p in picks if p["change_pct"] < 0]
    flat = [p for p in picks if p["change_pct"] == 0]
    avg_pct = round(sum(p["change_pct"] for p in picks) / len(picks), 2)

    by_tier = {}
    for p in picks:
        for tier in (p["tiers"].split(", ") if p["tiers"] else ["(untiered)"]):
            by_tier.setdefault(tier, []).append(p["change_pct"])
    tier_summary = {
        tier: {"count": len(vals), "avg_change_pct": round(sum(vals) / len(vals), 2)}
        for tier, vals in by_tier.items()
    }

    return {
        "total_tracked_picks": len(picks),
        "up": len(up), "down": len(down), "flat": len(flat),
        "win_rate_pct": round(len(up) / len(picks) * 100, 1),
        "average_change_pct": avg_pct,
        "by_tier": tier_summary,
        "picks": picks,
    }


def calibration_context() -> str:
    """The evaluation loop, closed.

    build_report() above measures how past picks ACTUALLY performed. On its
    own that's after-the-fact reporting. This turns the same real numbers
    into a compact factual block that agents/decision_agent.py injects into
    the next run's prompt — so the model is told how its previous calls
    turned out (overall, and per risk tier) before making new ones, and can
    calibrate its confidence against its own track record instead of
    starting from scratch every day.

    Returns "" when there's no tracked history yet — a cold start must not
    put an empty or misleading performance claim into a prompt. Never
    raises: a tracker problem should degrade the prompt, not break the run."""
    try:
        report = build_report()
    except Exception as e:  # noqa: BLE001
        print(f"[quality] calibration unavailable: {e}")
        return ""

    if not report.get("total_tracked_picks"):
        return ""

    lines = [
        f"Across {report['total_tracked_picks']} past pick(s) with real tracked price "
        f"history since they were listed: {report['up']} up, {report['down']} down, "
        f"{report['flat']} flat.",
        f"Realised win rate: {report['win_rate_pct']}%. "
        f"Average change since listing: {report['average_change_pct']:+.2f}%.",
    ]
    for tier, stats in (report.get("by_tier") or {}).items():
        lines.append(
            f"  - {tier}: {stats['count']} pick(s), average {stats['avg_change_pct']:+.2f}%"
        )
    return "\n".join(lines)


def print_report(report: dict):
    if not report.get("total_tracked_picks"):
        print(report.get("summary", "No data."))
        return
    print(f"Tracked picks with real since-listing history: {report['total_tracked_picks']}")
    print(f"  Up:   {report['up']}")
    print(f"  Down: {report['down']}")
    print(f"  Flat: {report['flat']}")
    print(f"  Win rate: {report['win_rate_pct']}%")
    print(f"  Average change since listing: {report['average_change_pct']:+.2f}%")
    print()
    print("By risk tier:")
    for tier, stats in report["by_tier"].items():
        print(f"  {tier}: {stats['count']} pick(s), avg {stats['avg_change_pct']:+.2f}%")


if __name__ == "__main__":
    report = build_report()
    print_report(report)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(report, indent=2, default=str))
    print(f"\nSaved full report -> {OUT_PATH}")
