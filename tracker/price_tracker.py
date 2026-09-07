"""
price_tracker.py  —  "The Ledger's Auditor"
================================================
A standalone tool, deliberately separate from agents/ — it isn't part of
the daily pipeline, it's a maintenance pass you run whenever you want to
check how past picks have actually performed.

    uv run python tracker/price_tracker.py

Walks every sheet in BOTH output/portfolio_tracker.xlsx (swing picks,
written by alerts/excel_report.py) and output/intraday_tracker.xlsx
(intraday picks, written by alerts/intraday_excel_report.py) — same
button, same tool, both workbooks. For every table it finds (anything
with a header row starting "Symbol"), it appends ONE new column dated
with today's date: each stock's real current price, plus its % change
vs. whichever column was the most recent price checkpoint for that table
— either the original "Current Price" column (first time this tool runs
against a sheet) or the last tracking column a previous run of this tool
added (every time after that).

Zero AI — this is a real price lookup and a subtraction, nothing more.
Every distinct symbol's price is fetched once per run and reused across
BOTH workbooks and every sheet/table in each, so a stock that shows up in
five different sheets only costs one real network call, not five.

Always appends — never overwrites a column, including on a same-day
re-run (the first run of the day gets a clean "2026-09-01" column; any
same-day re-run after that gets the actual "2026-09-01 14:32:07" it ran
at, instead of overwriting the earlier one). An earlier version overwrote
same-day columns in place, trusting a date-string match as proof it was a
safe re-run; that assumption broke in practice and
silently destroyed real tracked history, so this tool no longer deletes
anything, ever.
"""

import datetime as dt
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from data.market_data import get_market_stats  # noqa: E402

OUTPUT_DIR = ROOT / "output"
WORKBOOK_PATHS = [OUTPUT_DIR / "portfolio_tracker.xlsx", OUTPUT_DIR / "intraday_tracker.xlsx"]
# Back-compat: some callers/tests still reach for the single-file name.
WORKBOOK_PATH = WORKBOOK_PATHS[0]
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")  # prefix match — see _unique_header_label
_PRICE_RE = re.compile(r"₹?([\d,]+\.?\d*)")


def _find_tables(ws) -> list:
    """
    Locate every table in a sheet: a header row whose column-A cell is
    literally "Symbol", followed by data rows until the first blank one.
    """
    tables = []
    row, max_row = 1, ws.max_row
    while row <= max_row:
        if ws.cell(row=row, column=1).value == "Symbol":
            data_start = row + 1
            r = data_start
            while r <= max_row and ws.cell(row=r, column=1).value not in (None, ""):
                r += 1
            data_end = r - 1
            hc = 1
            while ws.cell(row=row, column=hc).value not in (None, ""):
                hc += 1
            last_col = hc - 1
            tables.append({"header_row": row, "data_start": data_start, "data_end": data_end, "last_col": last_col})
            row = data_end + 1
        else:
            row += 1
    return tables


def _reference_price_col(ws, table: dict):
    """The rightmost column in this table's header that's a real price
    checkpoint (the original "Current Price" column, or a previous
    tracking run's date column) — skips non-price columns in between,
    like Description or Tier(s)."""
    for col in range(table["last_col"], 0, -1):
        header = str(ws.cell(row=table["header_row"], column=col).value or "")
        if "Current Price" in header or _DATE_RE.match(header):
            return col
    return None


def _unique_header_label(ws, header_row: int, last_col: int, today: str) -> str:
    """`today` (just the date) for this table's first run of the day —
    keeps the common case (one run/day) reading clean. If this table
    already has a column labeled today, this is a same-day re-run: label
    the new column with the actual date+time of THIS run instead of
    overwriting the earlier one — more useful than a bare counter for
    telling same-day re-runs apart, and still guaranteed unique. A
    previous version overwrote same-day columns in place, trusting a
    date-string match as proof it was a harmless re-run; that assumption
    broke in practice and silently destroyed two real days of tracked
    data. This tool's only job is a historical record, so it now never
    deletes a column, full stop — worst case is a redundant same-day
    column, never a lost one."""
    existing = {
        str(ws.cell(row=header_row, column=c).value or "")
        for c in range(1, last_col + 1)
    }
    if today not in existing:
        return today
    label = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    n = 2
    base = label
    while label in existing:
        label = f"{base} ({n})"
        n += 1
    return label


def _parse_price(cell_value):
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    if isinstance(cell_value, str):
        m = _PRICE_RE.match(cell_value.strip())
        if m:
            return float(m.group(1).replace(",", ""))
    return None


def _real_price(symbol: str, cache: dict):
    if symbol not in cache:
        try:
            cache[symbol] = get_market_stats(symbol)["current_price"]
        except Exception as e:  # noqa: BLE001
            print(f"[tracker] {symbol}: could not fetch a real price ({e}) — skipped")
            cache[symbol] = None
    return cache[symbol]


def _track_workbook(path: Path, today: str, price_cache: dict) -> tuple:
    """Appends today's tracking column to every table in every sheet of
    ONE workbook. Returns (updated_cells, tables_seen) for that file."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    wb = load_workbook(path)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0CA30C")  # green, to visually mark tracker-added columns

    updated_cells = 0
    tables_seen = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for table in _find_tables(ws):
            # Always append — never overwrite an existing column. See
            # _unique_header_label's docstring for why: overwriting a
            # same-day column used to be "safe" by design, until it
            # silently destroyed two real days of tracked data.
            target_col = table["last_col"] + 1
            ref_col = _reference_price_col(ws, table)
            label = _unique_header_label(ws, table["header_row"], table["last_col"], today)
            tables_seen += 1

            header_cell = ws.cell(row=table["header_row"], column=target_col, value=label)
            header_cell.font, header_cell.fill = header_font, header_fill

            for row in range(table["data_start"], table["data_end"] + 1):
                symbol = ws.cell(row=row, column=1).value
                if not symbol:
                    continue
                current = _real_price(symbol, price_cache)
                if current is None:
                    ws.cell(row=row, column=target_col, value="n/a")
                    continue

                ref_price = _parse_price(ws.cell(row=row, column=ref_col).value) if ref_col else None
                if ref_price:
                    pct = (current - ref_price) / ref_price * 100
                    ws.cell(row=row, column=target_col, value=f"₹{current:,.2f} ({pct:+.1f}%)")
                else:
                    ws.cell(row=row, column=target_col, value=f"₹{current:,.2f}")
                updated_cells += 1

    wb.save(path)
    print(f"[tracker] {today}: updated {updated_cells} cells across "
          f"{tables_seen} table(s) in {len(wb.sheetnames)} sheet(s) -> {path}")
    return updated_cells, tables_seen


def run():
    existing = [p for p in WORKBOOK_PATHS if p.exists()]
    if not existing:
        raise RuntimeError(
            f"No workbook found at {WORKBOOK_PATHS[0]} or {WORKBOOK_PATHS[1]} — "
            "run run.py or intraday_run.py at least once first so there's something to track."
        )
    for missing in set(WORKBOOK_PATHS) - set(existing):
        print(f"[tracker] {missing.name} doesn't exist yet — skipping it")

    today = dt.date.today().isoformat()
    price_cache = {}  # shared across both workbooks — one real fetch per symbol, not per file
    for path in existing:
        _track_workbook(path, today, price_cache)


if __name__ == "__main__":
    run()
